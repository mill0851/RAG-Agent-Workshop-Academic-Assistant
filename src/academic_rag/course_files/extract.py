"""Turn Canvas HTML and downloaded or uploaded files into plain-text Segments,
keeping where each piece came from (PDF page, slide number)."""

from __future__ import annotations

import io
import logging
import re
from dataclasses import dataclass
from pathlib import PurePosixPath

from bs4 import BeautifulSoup

log = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".pptx", ".txt", ".md", ".html", ".htm", ".csv"}
# Bump when extraction improves so already-saved files get re-extracted on the next scrape.
EXTRACTOR_VERSION = 2  # 2: control characters removed

_CONTROL = re.compile(r"[\x00-\x08\x0e-\x1f\x7f]")  # all control characters except tab and line breaks


@dataclass
class Segment:
    """A unit of text with an optional location label (PDF page, slide number)."""

    text: str
    label: str = ""  # e.g. "p. 3" or "slide 5"


def html_to_text(html: str | None) -> str:
    """Convert Canvas rich-content HTML to readable text, keeping headings,
    list structure, and link targets (so 'see the rubric' keeps its URL)."""
    if not html or not html.strip():
        return ""
    soup = BeautifulSoup(html, "html.parser")
    for tag in soup(["script", "style", "noscript"]):
        tag.decompose()

    for a in soup.find_all("a"):
        href = a.get("href", "")
        if "verifier=" in href:  # a Canvas file link's access key: never save it
            href = href.split("?", 1)[0]
        label = a.get_text(" ", strip=True)
        if href and label and href.startswith(("http://", "https://")) and href != label:
            a.replace_with(f"{label} ({href})")
    for level in range(1, 7):
        for h in soup.find_all(f"h{level}"):
            h.insert_before("\n\n" + "#" * level + " ")
            h.append("\n\n")  # blank line, so the heading is its own paragraph
    for li in soup.find_all("li"):
        li.insert_before("\n- ")
    for br in soup.find_all("br"):
        br.replace_with("\n")
    for tag in soup.find_all(["p", "div", "tr", "table", "blockquote", "pre"]):
        tag.append("\n")

    text = soup.get_text()
    text = text.replace("\xa0", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r" *\n *", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_file(filename: str, data: bytes) -> list[Segment]:
    """Extract text from a downloaded or uploaded file. Returns [] when unsupported or empty."""
    ext = PurePosixPath(filename.lower()).suffix
    segments: list[Segment] = []
    try:
        if ext == ".pdf":
            segments = _pdf(data)
        elif ext == ".docx":
            segments = _docx(data)
        elif ext == ".pptx":
            segments = _pptx(data)
        elif ext in {".html", ".htm"}:
            segments = [Segment(html_to_text(data.decode("utf-8", "replace")))]
        elif ext in {".txt", ".md", ".csv"}:
            segments = [Segment(data.decode("utf-8", "replace"))]
    except Exception as exc:  # corrupt or encrypted files must not abort a whole ingest
        log.warning("could not extract %s: %s", filename, type(exc).__name__)
    cleaned = (Segment(_clean(s.text), s.label) for s in segments)
    return [s for s in cleaned if s.text]


def _clean(text: str) -> str:
    """PDFs often carry page breaks and font-control characters: keep the words, drop those."""
    return _CONTROL.sub("", text.replace("\x0c", "\n").replace("\x0b", "\n")).strip()


def _pdf(data: bytes) -> list[Segment]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception:
            return []
    segments = []
    for i, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            segments.append(Segment(text, f"p. {i}"))
    return segments


def _docx(data: bytes) -> list[Segment]:
    import docx

    doc = docx.Document(io.BytesIO(data))
    lines: list[str] = []
    for p in doc.paragraphs:
        if not p.text.strip():
            continue
        style = (p.style.name or "") if p.style is not None else ""
        m = re.match(r"Heading (\d)", style)
        lines.append(("#" * int(m.group(1)) + " " + p.text) if m else p.text)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                lines.append(" | ".join(cells))
    text = "\n\n".join(lines).strip()
    return [Segment(text)] if text else []


def _pptx(data: bytes) -> list[Segment]:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    segments = []
    for i, slide in enumerate(prs.slides, start=1):
        parts: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text.strip())
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    parts.append(" | ".join(c.text.strip() for c in row.cells))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append("Speaker notes: " + notes)
        text = "\n".join(p for p in parts if p).strip()
        if text:
            segments.append(Segment(text, f"slide {i}"))
    return segments
