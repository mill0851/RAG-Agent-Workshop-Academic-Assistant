"""A fake Canvas LMS served through respx, plus helpers to build real PDF/DOCX/PPTX bytes."""

from __future__ import annotations

import io
import time
from pathlib import Path

import httpx
import pytest
import respx
from fastapi.testclient import TestClient

from academic_rag.config import Settings
from academic_rag.index import HashEmbedder
from academic_rag.serve.api import create_app

BASE = "https://canvas.test.edu"
TOKEN = "TESTTOKEN-1234567890"


# ------------------------------------------------------------ file builders
def make_pdf(pages: list[str]) -> bytes:
    n = len(pages)
    kids = " ".join(f"{4 + 2 * i} 0 R" for i in range(n))
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        f"<< /Type /Pages /Kids [{kids}] /Count {n} >>".encode(),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    for i, text in enumerate(pages):
        objs.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {5 + 2 * i} 0 R "
            f"/Resources << /Font << /F1 3 0 R >> >> >>".encode()
        )
        stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
        objs.append(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
    out, offsets = b"%PDF-1.4\n", []
    for i, o in enumerate(objs, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + o + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n".encode() + b"0000000000 65535 f \n"
    out += b"".join(f"{o:010d} 00000 n \n".encode() for o in offsets)
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
    return out


def make_docx(heading: str, body: str) -> bytes:
    import docx

    d = docx.Document()
    d.add_heading(heading, level=1)
    d.add_paragraph(body)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def make_pptx(title: str, body: str) -> bytes:
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title
    slide.placeholders[1].text = body
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


# -------------------------------------------------------------- fake canvas
def _tab(tab_id: str, label: str | None = None) -> dict:
    return {"id": tab_id, "label": label or tab_id.title(), "type": "internal", "visibility": "public"}


class FakeCanvas:
    """Dispatches Canvas API requests. Mutate the dicts to simulate course edits."""

    def __init__(self) -> None:
        self.requests: list[httpx.Request] = []
        self.rate_limit_next = 0  # respond 403 'Rate Limit Exceeded' this many times
        self.files_enabled = {"101": True, "102": False}

        self.courses = {
            "101": {
                "id": 101, "name": "Intro to Machine Learning", "course_code": "CS 4731",
                "default_view": "wiki", "term": {"name": "Fall 2026"}, "public_description": "",
                "syllabus_body": "<h2>Grading</h2><p>Homework 40%, midterm 25%, final 35%. "
                                 "Late work loses 10% per day.</p>",
                "enrollments": [{"type": "student", "computed_current_score": 91.5,
                                 "computed_current_grade": "A-", "computed_final_score": 88.0,
                                 "computed_final_grade": "B+"}],
            },
            # Canvas name is just the numeric id; the real name has to come from the syllabus.
            "102": {
                "id": 102, "name": "93756976", "course_code": "CHEM 2053", "default_view": "modules",
                "term": {"name": "Fall 2026"},
                "syllabus_body": "<h1>Organic Chemistry I Syllabus</h1><p>This course covers the structure, "
                                 "bonding and reactivity of organic molecules, including nomenclature and "
                                 "stereochemistry. Office hours: Tuesdays 2-4pm in Room 305.</p>",
            },
        }
        self.tabs = {
            "101": [_tab(t) for t in ("home", "announcements", "assignments", "discussions", "grades", "people",
                                      "pages", "files", "syllabus", "quizzes", "modules")]
                   + [_tab("context_external_tool_9", "Piazza")],
            "102": [_tab(t) for t in ("home", "announcements", "syllabus", "files")],
        }
        self.front_page = {
            "101": {"page_id": 10, "url": "welcome", "title": "Welcome to ML 101", "front_page": True,
                    "updated_at": "2026-08-20T00:00:00Z", "html_url": f"{BASE}/courses/101/pages/welcome",
                    "body": "<p>Welcome! Start with Week 1 in Modules. Office hours are Mondays at 3pm.</p>"},
        }
        self.pages = {
            "101": {
                "gradient-descent": {
                    "page_id": 11, "url": "gradient-descent", "title": "Gradient Descent",
                    "updated_at": "2026-09-01T00:00:00Z",
                    "body": "<h1>Gradient Descent</h1><p>Gradient descent updates parameters in the "
                            "direction of the negative gradient scaled by the learning rate.</p>"
                            "<ul><li>A learning rate that is too large diverges</li></ul>",
                }
            },
            "102": {},
        }
        self.assignments = {
            "101": {
                5001: {
                    "id": 5001, "name": "Homework 2: Logistic Regression",
                    "html_url": f"{BASE}/courses/101/assignments/5001",
                    "due_at": "2026-10-20T04:59:00Z", "points_possible": 100,
                    "submission_types": ["online_upload"],
                    "description": "<p>Implement logistic regression from scratch using NumPy. "
                                   "Submit a notebook.</p>",
                    "rubric": [{"description": "Correctness", "long_description": "Gradients are <b>right</b>",
                                "points": 60, "ratings": [{"points": 60, "description": "Full marks"},
                                                          {"points": 0, "description": "No credit"}]}],
                    "submission": {"workflow_state": "graded", "score": 88.0, "grade": "88", "late": False,
                                   "missing": False, "excused": False, "submitted_at": "2026-10-19T20:00:00Z"},
                },
                5002: {  # no description at all: still needs to be findable by name and due date
                    "id": 5002, "name": "Project Proposal", "html_url": f"{BASE}/courses/101/assignments/5002",
                    "due_at": "2026-11-05T05:59:00Z", "points_possible": 20, "submission_types": ["online_text_entry"],
                    "description": "", "submission": {"workflow_state": "unsubmitted", "missing": False},
                },
            },
            "102": {},
        }
        self.quizzes = {
            "101": [{"id": 6001, "title": "Quiz 1: Linear Models", "html_url": f"{BASE}/courses/101/quizzes/6001",
                     "due_at": "2026-10-10T05:00:00Z", "points_possible": 10, "time_limit": 30,
                     "allowed_attempts": 2, "question_count": 8,
                     "description": "<p>Covers linear regression and the normal equations.</p>"}],
            "102": [],
        }
        self.announcements = {
            "101": [{"id": 7001, "title": "Midterm moved to Oct 28",
                     "html_url": f"{BASE}/courses/101/discussion_topics/7001",
                     "posted_at": "2026-10-01T12:00:00Z",
                     "message": "<p>The midterm exam is now on October 28 in Room 210.</p>"}],
            "102": [],
        }
        self.discussions = {
            "101": [
                {"id": 8001, "title": "Week 3 Q&A", "html_url": f"{BASE}/courses/101/discussion_topics/8001",
                 "message": "<p>Post your questions about regularization here.</p>", "is_announcement": False},
                {"id": 7001, "title": "Midterm moved to Oct 28", "is_announcement": True,
                 "message": "<p>duplicate of the announcement</p>"},
            ],
            "102": [],
        }
        self.folders = {
            "101": [{"id": 1, "full_name": "course files"}, {"id": 2, "full_name": "course files/Lectures"}],
            "102": [],
        }
        self.files = {
            "101": [
                self._file(9001, "lecture1.pdf", "application/pdf", 2),
                self._file(9002, "study_guide.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", 1),
                self._file(9003, "overfitting.pptx", "application/vnd.openxmlformats-officedocument.presentationml.presentation", 2),
                self._file(9004, "data.zip", "application/zip", 1),
                {**self._file(9005, "secret.pdf", "application/pdf", 1), "locked": True},
                self._file(9006, "reading_list.txt", "text/plain", 2),
                self._file(9007, "scanned.pdf", "application/pdf", 1),
            ],
            "102": [],
        }
        self.blobs = {
            "/files/9001/lecture1.pdf": make_pdf([
                "Lecture 1. Backpropagation uses the chain rule to compute gradients layer by layer.",
                "Regularization penalizes large weights using the L2 norm.",
            ]),
            "/files/9002/study_guide.docx": make_docx(
                "Study Guide", "Know the bias variance tradeoff and cross validation for the final exam."),
            "/files/9003/overfitting.pptx": make_pptx(
                "Overfitting", "Overfitting occurs when a model memorizes noise in the training data"),
            "/files/9004/data.zip": b"PK\x03\x04",
            "/files/9005/secret.pdf": make_pdf(["should never be indexed"]),
            "/files/9006/reading_list.txt": b"Required reading: chapter 3 of the deep learning textbook on convolutional networks.",
            "/files/9007/scanned.pdf": make_pdf([""]),  # no text layer
        }
        self.modules = {
            "101": [{"id": 1, "name": "Week 1: Optimization"}, {"id": 2, "name": "Week 2: Generalization"}],
            "102": [],
        }
        self.module_items = {
            "101": {
                1: [
                    {"type": "Page", "page_url": "gradient-descent", "title": "Gradient Descent"},
                    {"type": "Assignment", "content_id": 5001, "title": "Homework 2"},
                    {"type": "SubHeader", "title": "Readings"},
                    {"type": "ExternalUrl", "title": "Deep Learning Book", "external_url": "https://www.deeplearningbook.org"},
                ],
                2: [
                    {"type": "File", "content_id": 9001, "title": "lecture1.pdf"},
                    {"type": "File", "content_id": 9003, "title": "overfitting.pptx"},
                    {"type": "Quiz", "content_id": 6001, "title": "Quiz 1: Linear Models"},
                ],
            }
        }

    @staticmethod
    def _file(fid: int, name: str, ctype: str, folder_id: int) -> dict:
        return {
            "id": fid, "display_name": name, "filename": name, "content-type": ctype, "folder_id": folder_id,
            "size": 1000, "modified_at": "2026-09-10T10:00:00Z",
            "url": f"{BASE}/files/{fid}/download?download_frd=1&verifier=abc123",
            "_name": name,
        }

    # ---------------------------------------------------------------- routing
    def handle(self, request: httpx.Request) -> httpx.Response:
        self.requests.append(request)
        path, q = request.url.path, request.url.params

        if request.url.host == "cdn.test.edu":  # CDN: must never see the Canvas token
            blob = self.blobs.get(path)
            return httpx.Response(200, content=blob) if blob is not None else httpx.Response(404)

        if request.headers.get("authorization") != f"Bearer {TOKEN}":
            return httpx.Response(401, json={"errors": [{"message": "Invalid access token."}]})
        if self.rate_limit_next > 0:
            self.rate_limit_next -= 1
            return httpx.Response(403, text="403 Forbidden (Rate Limit Exceeded)")

        if path == "/api/v1/users/self":
            return httpx.Response(200, json={"id": 1, "name": "Test Student"})

        if path == "/api/v1/courses":
            if q.get("enrollment_state", "active") == "completed":
                return httpx.Response(200, json=[self.courses["101"]])
            visible = [self.courses["101"], {"id": 103, "access_restricted_by_date": True}]
            if q.get("page") == "2":
                return httpx.Response(200, json=[self.courses["102"]])
            return httpx.Response(
                200, json=visible,
                headers={"Link": f'<{BASE}/api/v1/courses?page=2&per_page=100&enrollment_state=active>; rel="next"'},
            )

        if path == "/api/v1/announcements":
            assert q.get("start_date"), "announcements need start_date or Canvas returns only 14 days"
            cid = q["context_codes[]"].removeprefix("course_")
            return httpx.Response(200, json=self.announcements.get(cid, []))

        parts = path.removeprefix("/api/v1/").lstrip("/").split("/")
        if parts[0] == "courses" and len(parts) >= 2:
            cid = parts[1]
            rest = parts[2:]
            if cid not in self.courses:
                return httpx.Response(404, json={"errors": [{"message": "not found"}]})
            if not rest:
                return httpx.Response(200, json=self.courses[cid])
            if rest == ["tabs"]:
                return httpx.Response(200, json=self.tabs[cid])
            if rest == ["front_page"]:
                page = self.front_page.get(cid)
                return httpx.Response(200, json=page) if page else httpx.Response(
                    404, json={"errors": [{"message": "page not found"}]})
            if rest == ["pages"]:
                stubs = [{k: v for k, v in p.items() if k != "body"} for p in self.pages[cid].values()]
                if cid in self.front_page:
                    stubs.insert(0, {k: v for k, v in self.front_page[cid].items() if k != "body"})
                return httpx.Response(200, json=stubs)
            if rest[0] == "pages":
                return httpx.Response(200, json=self.pages[cid][rest[1]])
            if rest == ["assignments"]:
                return httpx.Response(200, json=list(self.assignments[cid].values()))
            if rest == ["quizzes"]:
                return httpx.Response(200, json=self.quizzes[cid])
            if rest == ["discussion_topics"]:
                return httpx.Response(200, json=self.discussions[cid])
            if rest == ["folders"]:
                return httpx.Response(200, json=self.folders[cid])
            if rest == ["files"]:
                if not self.files_enabled[cid]:
                    return httpx.Response(403, json={"errors": [{"message": "user not authorized"}]})
                return httpx.Response(200, json=[{k: v for k, v in f.items() if k != "_name"} for f in self.files[cid]])
            if rest[0] == "files" and len(rest) == 2:  # one file, even when the Files tab is hidden
                f = next((f for f in self.files[cid] if str(f["id"]) == rest[1]), None)
                if f is None or f.get("locked"):
                    return httpx.Response(404 if f is None else 403, json={"errors": [{"message": "unauthorized"}]})
                return httpx.Response(200, json={k: v for k, v in f.items() if k != "_name"})
            if rest == ["modules"]:
                return httpx.Response(200, json=self.modules[cid])
            if rest[0] == "modules" and rest[2:] == ["items"]:
                return httpx.Response(200, json=self.module_items.get(cid, {}).get(int(rest[1]), []))

        if parts[0] == "files" and parts[2] == "download":  # Canvas bounces downloads to its CDN
            fid = int(parts[1])
            name = next(f["_name"] for fs in self.files.values() for f in fs if f["id"] == fid)
            return httpx.Response(302, headers={"Location": f"https://cdn.test.edu/files/{fid}/{name}"})

        return httpx.Response(404, json={"errors": [{"message": f"unhandled {path}"}]})

    def requested_paths(self) -> list[str]:
        return [r.url.path for r in self.requests]


# ----------------------------------------------------------------- fixtures
@pytest.fixture
def canvas():
    fake = FakeCanvas()
    with respx.mock(assert_all_called=False) as router:
        router.route().mock(side_effect=fake.handle)
        yield fake


@pytest.fixture
def settings(tmp_path: Path) -> Settings:
    return Settings(
        canvas_base_url=BASE, canvas_api_token=TOKEN, embedding_backend="hash",
        chunk_chars=500, chunk_overlap=60, data_dir=tmp_path / "data",
    )


@pytest.fixture(autouse=True)
def no_sleep(monkeypatch):
    monkeypatch.setattr("academic_rag.canvas.client.time.sleep", lambda s: None)


# ---------------------------------------------------------------- API client
@pytest.fixture
def client(canvas, settings):
    return TestClient(create_app(settings, embedder=HashEmbedder()))


def wait(client, job, timeout=30):
    deadline = time.time() + timeout
    while time.time() < deadline:
        view = client.get(f"/jobs/{job['id']}").json()
        if view["status"] in {"done", "error"}:
            return view
        time.sleep(0.05)
    raise AssertionError("job did not finish")


@pytest.fixture
def loaded(client):
    """A client whose server has already scraped and embedded the fake Canvas."""
    for path in ("/scrape", "/embed"):
        r = client.post(path)
        assert r.status_code == 202
        assert wait(client, r.json())["status"] == "done"
    return client


# ---------------------------------------------------------------- a fake AI
class FakeProvider:
    """Stands in for a real model: remembers what it was asked, calls the tools it is told
    to (the way a real model would), then answers in fixed pieces."""

    name = "fake"
    model = ""

    def __init__(self, pieces=("The midterm is on October 28 ", "[1]."), tool_calls=()):
        self.pieces = pieces
        self.tool_calls = list(tool_calls)  # [(tool name, arguments)] to call before answering
        self.calls = []  # (system, prompt, effort) for each question
        self.results = []  # what each tool call returned
        self.mcp_urls = []  # where Claude Code or Codex would have reached the tools

    def answer(self, system, prompt, tools, effort="medium"):
        self.calls.append((system, prompt, effort))
        self.mcp_urls.append(tools.mcp_url)
        for name, args in self.tool_calls:
            yield {"type": "tool", "name": name, "args": args}
            self.results.append(tools.call(name, args))
        for piece in self.pieces:
            yield {"type": "text", "text": piece}
